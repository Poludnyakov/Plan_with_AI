"""Real binary files exercise local readers; external OCR/LLM calls are mocked."""
import asyncio
import io
import json
import re
import struct
import threading
import zipfile
from contextlib import AsyncExitStack
from datetime import date, datetime, time, timedelta
from unittest.mock import AsyncMock, patch
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject, DictionaryObject, NameObject, NumberObject,
)
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

import schedule_document_service as service
import spreadsheet_schedule_parser as spreadsheet_parser
from database import Base
from schedule_models import ScheduleSeries
from schedule_service import confirm_import


@pytest.fixture
def anyio_backend():
    return "asyncio"


def xlsx_bytes(build):
    workbook = Workbook()
    build(workbook)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def rewrite_zip(content, changes):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            name, payload = info.filename, source.read(info)
            if name in changes:
                name, payload = changes[name](name, payload)
            target.writestr(name, payload)
    return output.getvalue()


def pdf_bytes(kinds, password=None):
    writer = PdfWriter()
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    font_ref = writer._add_object(font)
    for index, kind in enumerate(kinds, 1):
        page = writer.add_blank_page(width=612, height=792)
        if kind == "blank":
            continue
        resources = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})})
        stream = DecodedStreamObject()
        if kind == "scan":
            image = DecodedStreamObject()
            image.set_data(b"\xff\xff\xff")
            image.update({
                NameObject("/Type"): NameObject("/XObject"),
                NameObject("/Subtype"): NameObject("/Image"),
                NameObject("/Width"): NumberObject(1),
                NameObject("/Height"): NumberObject(1),
                NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
                NameObject("/BitsPerComponent"): NumberObject(8),
            })
            resources[NameObject("/XObject")] = DictionaryObject({NameObject("/Im1"): writer._add_object(image)})
            stream.set_data(b"q 600 0 0 780 0 0 cm /Im1 Do Q")
        elif kind == "rotated":
            stream.set_data(b"BT /F1 12 Tf 0 1 -1 0 100 100 Tm (Rotated Monday Training 09:00-10:00) Tj ET")
        elif kind == "table":
            stream.set_data(
                b"BT /F1 12 Tf 20 700 Td (Day) Tj ET\n"
                b"BT /F1 12 Tf 180 700 Td (Time) Tj ET\n"
                b"BT /F1 12 Tf 360 700 Td (Subject) Tj ET\n"
                b"BT /F1 12 Tf 20 680 Td (Monday) Tj ET\n"
                b"BT /F1 12 Tf 180 680 Td (09:00-10:00) Tj ET\n"
                b"BT /F1 12 Tf 360 680 Td (Training) Tj ET"
            )
        else:
            label = "Hi" if kind == "short" else f"Page{index} Thursday Training 03.09.2026 09:00-10:00"
            stream.set_data(f"BT /F1 12 Tf 20 700 Td ({label}) Tj ET".encode())
        page[NameObject("/Resources")] = resources
        page[NameObject("/Contents")] = writer._add_object(stream)
    if password is not None:
        writer.encrypt(password)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


@pytest.mark.anyio
@pytest.mark.parametrize("dimension", [b"A1:A1", b"A1:XFD1048576", None])
async def test_xlsx_uses_actual_cells_instead_of_reported_dimensions(dimension):
    content = xlsx_bytes(lambda wb: wb.active.append(["Monday", "09:00-10:00", "Training"]))
    content = rewrite_zip(content, {"xl/worksheets/sheet1.xml": lambda name, data: (
        name, re.sub(br'<dimension[^>]*/>', b'' if dimension is None else b'<dimension ref="' + dimension + b'"/>', data)
    )})
    text = await service.extract_document_text(content, "test.xlsx")
    assert "Monday\t09:00-10:00\tTraining" in text


@pytest.mark.anyio
async def test_merge_ranges_follow_relationship_paths_and_sheet_order():
    def build(wb):
        wb.active.title = "First"
        wb.active["A1"] = "Alpha"
        wb.active.merge_cells("A1:C1")
        second = wb.create_sheet("Second")
        second["B1"] = "Beta"
        second.merge_cells("B1:D1")
    content = xlsx_bytes(build)
    content = rewrite_zip(content, {
        "xl/worksheets/sheet1.xml": lambda _, data: ("xl/worksheets/custom-20.xml", data),
        "xl/worksheets/sheet2.xml": lambda _, data: ("xl/worksheets/custom-3.xml", data),
        "xl/_rels/workbook.xml.rels": lambda name, data: (
            name, data.replace(b"sheet1.xml", b"custom-20.xml").replace(b"sheet2.xml", b"custom-3.xml")
        ),
    })
    text = await service.extract_document_text(content, "test.xlsx")
    assert "R1\tAlpha\tAlpha\tAlpha" in text
    assert "R1\t\tBeta\tBeta\tBeta" in text
    assert "ОБЪЕДИНЕНО B1:D1=Beta" in text


@pytest.mark.anyio
async def test_xlsx_dates_times_numbers_and_false_are_preserved():
    content = xlsx_bytes(lambda wb: wb.active.append([date(2026, 9, 3), time(9, 30), 0, False, 1.25]))
    text = await service.extract_document_text(content, "test.xlsx")
    assert "2026-09-03 00:00:00\t09:30:00\t0\tFalse\t1.25" in text


@pytest.mark.anyio
async def test_date_formula_chains_are_resolved_without_a_cached_result():
    def build(wb):
        wb.active["A1"] = "=DATE(2026,12,31)"
        wb.active["B1"] = "=$A$1+1"
        wb.active["C1"] = "=B1+1"
        wb.active["A3"] = "Training 09:00-10:00"
        wb.active["B3"] = "Science 11:00-12:00"
    text = await service.extract_document_text(xlsx_bytes(build), "test.xlsx")
    assert "R1\t2026-12-31\t2027-01-01\t2027-01-02" in text
    assert "КАЛЕНДАРНАЯ_ДАТА 2027-01-01 | ЯЧЕЙКА B3: Science" in text


@pytest.mark.anyio
async def test_formula_result_is_recomputed_instead_of_using_a_stale_date():
    def build(wb):
        wb.active["A1"] = datetime(2026, 12, 31)
        wb.active["B1"] = "=A1+1"
        wb.active["B1"].number_format = "yyyy-mm-dd"
    content = xlsx_bytes(build)
    content = rewrite_zip(content, {"xl/worksheets/sheet1.xml": lambda name, data: (
        name, data.replace(b"<f>A1+1</f><v></v>", b"<f>A1+1</f><v>43831</v>")
    )})
    text = await service.extract_document_text(content, "test.xlsx")
    assert "2027-01-01" in text
    assert "2020-01-01" not in text


@pytest.mark.anyio
@pytest.mark.parametrize("formula,message", [("=SUM(A1:A5)", "нет сохранённого результата"), ("=A1+1", "циклическая")])
async def test_unresolvable_formulas_fail_with_actionable_error(formula, message):
    content = xlsx_bytes(lambda wb: setattr(wb.active["A1"], "value", formula))
    with pytest.raises(ValueError, match=message):
        await service.extract_document_text(content, "test.xlsx")


@pytest.mark.anyio
async def test_excel_errors_do_not_become_events():
    content = xlsx_bytes(lambda wb: setattr(wb.active["A1"], "value", "#REF!"))
    with pytest.raises(ValueError, match="ошибка Excel"):
        await service.extract_document_text(content, "test.xlsx")


@pytest.mark.anyio
@pytest.mark.parametrize("build", [lambda wb: None, lambda wb: wb.create_sheet("Empty")])
async def test_empty_workbook_is_rejected(build):
    with pytest.raises(ValueError, match="непустых ячеек"):
        await service.extract_document_text(xlsx_bytes(build), "test.xlsx")


@pytest.mark.anyio
async def test_sparse_huge_cell_and_merge_are_rejected_before_iteration():
    for build in (
        lambda wb: setattr(wb.active["XFD1000"], "value", "Training"),
        lambda wb: wb.active.merge_cells("A1:A100"),
    ):
        with patch.object(service, "MAX_TABLE_CELLS", 50), pytest.raises(ValueError, match="слишком большой"):
            await service.extract_document_text(xlsx_bytes(build), "test.xlsx")


@pytest.mark.anyio
async def test_oversized_archive_and_sheet_count_are_rejected():
    content = xlsx_bytes(lambda wb: wb.create_sheet("Second"))
    with patch.object(service, "MAX_ARCHIVE_UNCOMPRESSED_BYTES", 100), pytest.raises(ValueError, match="Распакованный"):
        await service.extract_document_text(content, "test.xlsx")
    with patch.object(service, "MAX_SHEETS", 1), pytest.raises(ValueError, match="листов"):
        await service.extract_document_text(content, "test.xlsx")


@pytest.mark.anyio
async def test_xlsx_xml_entities_are_rejected():
    content = xlsx_bytes(lambda wb: wb.active.append(["Training"]))
    content = rewrite_zip(content, {"xl/worksheets/sheet1.xml": lambda name, data: (
        name, b'<!DOCTYPE worksheet [<!ENTITY attack "expanded">]>' + data.replace(b"Training", b"&attack;")
    )})
    with pytest.raises(ValueError, match="повреждён"):
        await service.extract_document_text(content, "test.xlsx")


@pytest.mark.anyio
@pytest.mark.parametrize("filename,content", [("bad.xlsx", b"broken zip"), ("bad.pdf", b"%PDF-1.7\nbroken")])
async def test_corrupt_binary_files_return_readable_errors(filename, content):
    with pytest.raises(ValueError, match="повреждён"):
        await service.extract_document_text(content, filename)


@pytest.mark.anyio
async def test_binary_signature_overrides_wrong_messenger_extension():
    content = xlsx_bytes(lambda wb: wb.active.append(["Training"]))
    assert "Training" in await service.extract_document_text(content, "document.pdf")
    assert "Training" in await service.extract_document_text(pdf_bytes(["text"]), "document.xlsx")


@pytest.mark.anyio
async def test_explicit_unsupported_extension_is_not_decoded_as_text():
    with pytest.raises(ValueError, match="Поддерживаются"):
        await service.extract_document_text(b"arbitrary data", "file.exe")


@pytest.mark.anyio
async def test_long_cell_and_total_text_are_not_silently_truncated():
    content = xlsx_bytes(lambda wb: wb.active.append(["x" * 6000]))
    assert "x" * 6000 in await service.extract_document_text(content, "test.xlsx")
    with patch.object(service, "MAX_DOCUMENT_CHARS", 100), pytest.raises(ValueError, match="слишком большой"):
        await service.extract_document_text(content, "test.xlsx")


@pytest.mark.anyio
async def test_short_and_blank_native_pdf_pages_do_not_request_ocr():
    with patch.object(service, "_vision_ocr", new_callable=AsyncMock) as ocr:
        text = await service.extract_document_text(pdf_bytes(["text", "short", "blank"]), "test.pdf")
    assert "Page1" in text and "Hi" in text
    assert "=== СТРАНИЦА 2 ===" in text
    ocr.assert_not_awaited()


@pytest.mark.anyio
async def test_pdf_table_layout_and_rotated_native_text_are_preserved():
    text = await service.extract_document_text(pdf_bytes(["table", "rotated"]), "test.pdf")
    row = next(line for line in text.splitlines() if "Monday" in line and "Training" in line and "Rotated" not in line)
    assert row.index("Monday") < row.index("09:00") < row.index("Training")
    assert "Rotated Monday Training" in text


@pytest.mark.anyio
async def test_hybrid_pdf_ocr_reads_only_scans_even_after_eighth_page():
    content = pdf_bytes(["text"] * 9 + ["scan", "text", "scan"])
    with patch.object(service.settings, "OPENROUTER_API_KEY", "test"), \
         patch.object(service, "_pdf_page_images", return_value=[b"page10", b"page12"]) as raster, \
         patch.object(service, "_vision_ocr", new_callable=AsyncMock, side_effect=["=== OCR СТРАНИЦА 1 ===\nScan10", "Scan12"]) as ocr:
        text = await service.extract_document_text(content, "test.pdf")
    raster.assert_called_once_with(content, [10, 12])
    assert [call.args for call in ocr.await_args_list] == [(b"page10", ".png"), (b"page12", ".png")]
    assert "Page9" in text and "Page11" in text
    assert text.index("Page9") < text.index("=== OCR СТРАНИЦА 10 ===") < text.index("Page11") < text.index("=== OCR СТРАНИЦА 12 ===")


@pytest.mark.anyio
async def test_scanned_pdf_over_budget_is_rejected_before_any_ocr():
    with patch.object(service, "_vision_ocr", new_callable=AsyncMock) as ocr, pytest.raises(ValueError, match="больше 8"):
        await service.extract_document_text(pdf_bytes(["scan"] * 9), "test.pdf")
    ocr.assert_not_awaited()


@pytest.mark.anyio
async def test_blank_pdf_is_not_mistaken_for_text():
    with pytest.raises(ValueError, match="не найден текст"):
        await service.extract_document_text(pdf_bytes(["blank"]), "test.pdf")


@pytest.mark.anyio
async def test_password_protected_pdf_and_empty_password_pdf():
    with pytest.raises(ValueError, match="защищён паролем"):
        await service.extract_document_text(pdf_bytes(["text"], "secret"), "test.pdf")
    assert "Training" in await service.extract_document_text(pdf_bytes(["text"], ""), "test.pdf")


@pytest.mark.anyio
async def test_pdf_page_stream_and_page_count_limits():
    with patch.object(service, "MAX_PDF_PAGE_STREAM_BYTES", 10), pytest.raises(ValueError, match="слишком большое"):
        await service.extract_document_text(pdf_bytes(["text"]), "test.pdf")
    with pytest.raises(ValueError, match="больше 200"):
        await service.extract_document_text(pdf_bytes(["blank"] * 201), "test.pdf")


@pytest.mark.anyio
async def test_ocr_failure_does_not_return_partial_native_document():
    with patch.object(service.settings, "OPENROUTER_API_KEY", "test"), \
         patch.object(service, "_pdf_page_images", return_value=[b"scan"]), \
         patch.object(service, "_vision_ocr", new_callable=AsyncMock, side_effect=ValueError("OCR failed")), \
         pytest.raises(ValueError, match="OCR failed"):
        await service.extract_document_text(pdf_bytes(["text", "scan"]), "test.pdf")


@pytest.mark.anyio
async def test_real_pdf_rasterizer_renders_requested_pages_at_bounded_resolution():
    images = await asyncio.to_thread(service._pdf_page_images, pdf_bytes(["text", "scan"]), [2])
    assert len(images) == 1
    assert images[0].startswith(b"\x89PNG\r\n\x1a\n")
    width, height = struct.unpack(">II", images[0][16:24])
    assert max(width, height) == 2000


@pytest.mark.anyio
async def test_native_pdf_pages_are_packed_without_loss():
    text = await service.extract_document_text(pdf_bytes(["text"] * 12), "test.pdf")
    chunks = service.document_chunks(text, max_chars=service.PDF_CHUNK_CHARS)
    assert len(chunks) == 1
    assert all(f"=== СТРАНИЦА {number} ===" in chunks[0] for number in range(1, 13))


@pytest.mark.anyio
async def test_local_reader_does_not_block_async_bot_loop():
    started, release = threading.Event(), threading.Event()
    content = xlsx_bytes(lambda wb: wb.active.append(["Training"]))
    def read(_):
        started.set()
        assert release.wait(2)
        return "Monday Training 09:00-10:00"
    with patch.object(service, "_xlsx_text", side_effect=read):
        task = asyncio.create_task(service.extract_document_text(content, "test.xlsx"))
        try:
            for _ in range(100):
                if started.is_set() or task.done():
                    break
                await asyncio.sleep(0.005)
            assert started.is_set()
            assert not task.done()
        finally:
            release.set()
        assert "Training" in await task


@pytest.mark.anyio
async def test_two_date_columns_in_record_are_not_treated_as_calendar_grid():
    content = xlsx_bytes(lambda wb: wb.active.append([date(2026, 9, 3), date(2026, 9, 4), "Training"]))
    text = await service.extract_document_text(content, "test.xlsx")
    assert service.CALENDAR_EVIDENCE_PREFIX not in text


@pytest.mark.anyio
async def test_llm_truncation_cannot_be_accepted_as_complete_json():
    response = SimpleNamespace(choices=[SimpleNamespace(
        message=SimpleNamespace(content='{"slots":[]}'), finish_reason="length",
    )])
    with patch.object(service, "chat_completion", new_callable=AsyncMock, return_value=response), \
         pytest.raises(ValueError, match="усечён"):
        await service._model_json("model", [])


def test_malformed_llm_rows_and_invalid_status_block_production_import():
    for payload in (
        {"slots": [{"title": "Training", "weekday": 0, "start_time": "11:00", "end_time": "09:00"}]},
        {"slots": {}}, {"status": "made_up", "slots": []},
    ):
        with pytest.raises(ValueError):
            service._validated_chunk(json.dumps(payload))


@pytest.mark.parametrize("clock,expected", [("23:55", ("23:55", "23:56")), ("23:59", ("23:58", "23:59")), ("09:00", ("09:00", "09:01"))])
def test_explicit_point_deadlines_have_a_short_valid_interval(clock, expected):
    parsed = service._validated_chunk(json.dumps({"slots": [{
        "title": "Дедлайн сдачи д/з", "occurrence_date": "2026-09-03",
        "start_time": clock, "end_time": clock, "source_quote": "Дедлайн " + clock,
    }]}))
    assert (parsed.slots[0]["start_time"], parsed.slots[0]["end_time"]) == expected


def test_explicit_invalid_deadline_interval_is_not_reinterpreted_as_a_point():
    with pytest.raises(ValueError):
        service._validated_chunk(json.dumps({"slots": [{
            "title": "Срок сдачи", "weekday": 0, "start_time": "23:00", "end_time": "22:00",
            "source_quote": "Срок сдачи 23:00-22:00",
        }]}))


def test_deadline_quote_may_also_mention_a_separate_lesson_time():
    parsed = service._validated_chunk(json.dumps({"slots": [{
        "title": "Дедлайн", "occurrence_date": "2026-09-03", "start_time": "23:55", "end_time": "23:55",
        "source_quote": "Вебинар 19:00. Дедлайн 23:55.",
    }]}))
    assert parsed.slots[0]["end_time"] == "23:56"


def test_model_null_optional_fields_and_russian_weekday_are_normalized():
    parsed = service._validated_chunk(json.dumps({"slots": [{
        "title": "Математика", "weekday": "понедельник", "start_time": "09:00", "end_time": "10:00",
        "description": None, "confidence": None, "week_pattern": None, "source_quote": None,
    }]}))
    assert parsed.slots[0]["weekday"] == 0
    assert parsed.slots[0]["description"] == ""


@pytest.mark.anyio
async def test_pdf_clarification_does_not_trigger_full_ocr_fallback():
    content = pdf_bytes(["text"])
    with patch.object(service.settings, "OPENROUTER_API_KEY", "test"), \
         patch.object(service, "_model_json", new_callable=AsyncMock, return_value='{"status":"clarification","clarification":"Choose group","choices":["A","B"]}'), \
         patch.object(service, "_vision_ocr", new_callable=AsyncMock) as ocr, \
         pytest.raises(service.ScheduleDocumentClarification, match="Choose group"):
        await service.parse_schedule_document(content, "test.pdf", "")
    ocr.assert_not_awaited()


@pytest.mark.anyio
async def test_pdf_cannot_import_one_page_while_another_page_needs_a_choice():
    row = {"title": "Training", "occurrence_date": "2026-09-03", "start_time": "09:00", "end_time": "10:00"}
    with patch.object(service.settings, "OPENROUTER_API_KEY", "test"), \
         patch.object(service, "document_chunks", return_value=["page1", "page2"]), \
         patch.object(service, "_model_json", new_callable=AsyncMock, side_effect=[
             json.dumps({"slots": [row]}), '{"status":"clarification","clarification":"Choose group","choices":["A","B"]}',
         ]), pytest.raises(service.ScheduleDocumentClarification, match="Choose group"):
        await service.parse_schedule_document(pdf_bytes(["text", "text"]), "test.pdf", "")


def test_title_verification_checks_the_complete_subject():
    accepted, rejected = service._verify_source_support(
        [{"title": "Урок математики"}, {"title": "Математика биология"}], "Математика 09:00-10:00",
    )
    assert [slot["title"] for slot in accepted] == ["Урок математики"]
    assert rejected == ["Математика биология"]


@pytest.mark.anyio
@pytest.mark.parametrize("platform", ["telegram", "max"])
@pytest.mark.parametrize("extension", [".xlsx", ".pdf"])
async def test_real_file_to_database_round_trip_for_both_messengers(platform, extension):
    # Real local readers, domain validation, draft creation and persistence.
    # Only the remote semantic model is replaced by a deterministic response.
    if extension == ".xlsx":
        content = xlsx_bytes(lambda wb: wb.active.append([date(2026, 9, 3), "09:00-10:00", "Training"]))
    else:
        content = pdf_bytes(["text"])
    row = {"title": "Training", "occurrence_date": "2026-09-03", "start_time": "09:00", "end_time": "10:00", "source_quote": "Training 09:00-10:00"}
    xlsx_answer = json.dumps({"status": "ok", "events": [
        {"cell": "C1", "title": "Training", "date": "2026-09-03", "start": "09:00", "end": "10:00", "kind": "event"},
    ]})
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    try:
        async with engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        async with AsyncExitStack() as stack:
            db = await stack.enter_async_context(sessions())
            stack.enter_context(patch.object(service.settings, "OPENROUTER_API_KEY", "test"))
            stack.enter_context(patch.object(service, "_model_json", new_callable=AsyncMock, return_value=json.dumps({"slots": [row]})))
            stack.enter_context(patch.object(spreadsheet_parser.settings, "OPENROUTER_MODEL_FOR_FILES", "test/model"))
            stack.enter_context(patch.object(spreadsheet_parser, "_ask_model", new_callable=AsyncMock, return_value=xlsx_answer))
            if platform == "telegram":
                from handlers.pipeline_handlers_intervals import create_verified_schedule_draft
                draft, extraction = await create_verified_schedule_draft(db, 123, content, "schedule" + extension, "")
            else:
                from max_bot.handler import MaxUpdateHandler
                anonymizer = SimpleNamespace(anonymize_text=lambda value: value, clean_event_title=lambda value: value, clean_display_text=lambda value: value)
                handler = MaxUpdateHandler(SimpleNamespace(), SimpleNamespace(anonymizer=anonymizer), SimpleNamespace(miniapp_name="test"))
                draft, extraction = await handler.create_verified_schedule_draft(db, 123, content, "schedule" + extension, "")
            assert len(extraction["slots"]) == 1
            assert draft.status == "ready"
            assert await confirm_import(db, platform, 123, draft.id) == ("created", 1)
            await confirm_import(db, platform, 123, draft.id)
            stored = (await db.execute(select(ScheduleSeries))).scalars().all()
            assert len(stored) == 1
            assert draft.source_platform == platform
            assert stored[0].account_id == draft.account_id
            assert stored[0].title == "Training"
            assert stored[0].valid_from == stored[0].valid_until == date(2026, 9, 3)
    finally:
        await engine.dispose()
