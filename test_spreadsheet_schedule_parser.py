"""Real XLSX files through the spreadsheet reader; only the remote model is faked."""
import asyncio
import io
import json
import re
from datetime import date, time, timedelta
from unittest.mock import patch

import pytest
from openpyxl import Workbook

import spreadsheet_schedule_parser as parser
from schedule_document_service import ScheduleDocumentClarification, parse_schedule_document


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.fixture(autouse=True)
def configured_model():
    with patch.object(parser.settings, "OPENROUTER_API_KEY", "test"), \
         patch.object(parser.settings, "OPENROUTER_MODEL_FOR_FILES", "test/model"):
        yield


def xlsx_bytes(build) -> bytes:
    workbook = Workbook()
    build(workbook)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


CELL_LINE = re.compile(r"^  ([A-Z]+[0-9]+)(?: \{(\d{4}-\d{2}-\d{2}) ..\}| \{..\})?(?: \(объединено [^)]+\))?: (.*)$")


def fragment_cells(user_content: str) -> list[tuple[str, str | None, str]]:
    """Cells the app asked the model to extract from (not the context)."""
    fragment = user_content.split("\nФРАГМЕНТ:", 1)[1]
    return [match.groups() for line in fragment.splitlines() if (match := CELL_LINE.match(line))]


def clock_answer(user_content: str) -> str:
    """A faithful fake model: one event per fragment cell that names a clock."""
    events = []
    for cell, hinted, text in fragment_cells(user_content):
        clock = re.search(r"(\d{1,2}):(\d{2})", text)
        if not clock:
            continue
        events.append({
            "cell": cell, "title": re.split(r" в | / |\s*\d{1,2}:\d{2}", text)[0], "date": hinted,
            "start": f"{int(clock[1]):02d}:{clock[2]}", "end": None,
            "kind": "deadline" if "Дедлайн" in text else "event",
        })
    return json.dumps({"status": "ok", "events": events}, ensure_ascii=False)


def run(content: bytes, prompt: str = "", **kwargs):
    return parse_schedule_document(content, "schedule.xlsx", prompt, **kwargs)


@pytest.mark.anyio
async def test_calendar_grid_gets_exact_dates_and_skips_cells_without_time():
    def build(wb):
        ws = wb.active
        ws.append(["ПН", "ВТ", "СР"])
        ws.append([date(2026, 9, 7), date(2026, 9, 8), date(2026, 9, 9)])
        ws.append(["Практика в 17:00 по мск.", "Дедлайн д/з в 23:55", "Выход нового урока"])
    seen = []

    async def model(user_content):
        seen.append(user_content)
        return clock_answer(user_content)

    with patch.object(parser, "_ask_model", side_effect=model):
        result = await run(xlsx_bytes(build))

    assert "A3 {2026-09-07 пн}: Практика в 17:00 по мск." in seen[0]
    assert [(s["occurrence_date"], s["start_time"], s["end_time"], s["title"]) for s in result["slots"]] == [
        ("2026-09-07", "17:00", "18:30", "Практика"),
        ("2026-09-08", "23:25", "23:55", "Дедлайн д/з"),
    ]
    assert result["requires_range"] is False
    assert result["pipeline"] == "spreadsheet_ai"
    assert "Выход нового урока" in result["verification"]["warnings"][0]


@pytest.mark.anyio
async def test_cell_date_from_the_grid_wins_over_a_wrong_model_date():
    def build(wb):
        wb.active.append([date(2026, 9, 7), date(2026, 9, 8)])
        wb.active.append(["Йога 09:00", "Бассейн 19:00"])
    answer = json.dumps({"events": [
        {"cell": "B2", "title": "Бассейн", "date": "2026-09-07", "start": "19:00", "end": None},
    ]})
    with patch.object(parser, "_ask_model", return_value=answer), \
         patch.object(parser, "_repair_missed", return_value=[]):
        result = await run(xlsx_bytes(build))
    assert [(s["title"], s["occurrence_date"]) for s in result["slots"]] == [("Бассейн", "2026-09-08")]


@pytest.mark.anyio
async def test_invented_time_title_and_cells_are_rejected():
    def build(wb):
        wb.active.append([date(2026, 9, 7), date(2026, 9, 8)])
        wb.active.append(["Йога 09:00", "Выход урока"])
    answer = json.dumps({"events": [
        {"cell": "A2", "title": "Йога", "start": "09:00"},
        {"cell": "B2", "title": "Выход урока", "start": "00:00"},
        {"cell": "A2", "title": "Футбол", "start": "09:00"},
        {"cell": "Z99", "title": "Йога", "start": "09:00"},
    ]})
    with patch.object(parser, "_ask_model", return_value=answer):
        result = await run(xlsx_bytes(build))
    assert [s["title"] for s in result["slots"]] == ["Йога"]
    assert "Отброшено неподтверждённых таблицей строк: 3." in result["verification"]["warnings"]


@pytest.mark.anyio
async def test_list_table_and_explicit_end_time():
    def build(wb):
        wb.active.append(["Дата", "Начало", "Конец", "Предмет"])
        wb.active.append([date(2026, 10, 12), time(9, 0), time(10, 30), "Математика"])
        wb.active.append(["13.10.2026", "14:00", None, "Химия"])
    answer = json.dumps({"events": [
        {"cell": "D2", "title": "Математика", "date": "2026-10-12", "start": "09:00", "end": "10:30"},
        {"cell": "D3", "title": "Химия", "date": "2026-10-13", "start": "14:00", "end": None},
    ]})
    with patch.object(parser, "_ask_model", return_value=answer):
        result = await run(xlsx_bytes(build))
    assert [(s["occurrence_date"], s["start_time"], s["end_time"]) for s in result["slots"]] == [
        ("2026-10-12", "09:00", "10:30"), ("2026-10-13", "14:00", "15:30"),
    ]


@pytest.mark.anyio
async def test_weekly_timetable_is_recurring_and_keeps_week_parity():
    def build(wb):
        wb.active.append(["День", "Время", "Предмет"])
        wb.active.append(["Понедельник", "08:30-10:00", "Алгебра (чёт.)"])
        wb.active.append(["Вторник", "10:10-11:40", "История"])
    answer = json.dumps({"events": [
        {"cell": "C2", "title": "Алгебра", "weekday": 0, "start": "08:30", "end": "10:00", "week": "even"},
        {"cell": "C3", "title": "История", "weekday": "1", "start": "10:10", "end": "11:40"},
    ]})
    with patch.object(parser, "_ask_model", return_value=answer):
        result = await run(xlsx_bytes(build))
    assert [(s["weekday"], s["occurrence_date"], s["week_pattern"]) for s in result["slots"]] == [
        (0, None, "even"), (1, None, "every"),
    ]
    assert result["requires_range"] is True


@pytest.mark.anyio
async def test_large_annual_calendar_is_split_read_in_parallel_and_complete():
    def build(wb):
        for week in range(31):
            wb.active.append([date(2026, 9, 7) + timedelta(days=week * 7 + day) for day in range(7)])
            wb.active.append([f"Урок {week}-{day} в 09:00" for day in range(7)])
    active = peak = 0

    async def model(user_content):
        nonlocal active, peak
        active += 1
        peak = max(peak, active)
        await asyncio.sleep(0.01)
        active -= 1
        cells = fragment_cells(user_content)
        assert len(cells) <= parser.CHUNK_CELLS
        return clock_answer(user_content)

    with patch.object(parser, "_ask_model", side_effect=model), \
         patch.object(parser.settings, "OPENROUTER_DOCUMENT_CONCURRENCY", 3):
        result = await run(xlsx_bytes(build))
    assert len(result["slots"]) == 217
    assert len({s["occurrence_date"] for s in result["slots"]}) == 217
    assert result["chunks_processed"] > 1
    assert peak == 3


@pytest.mark.anyio
async def test_events_from_context_rows_are_not_duplicated():
    def build(wb):
        wb.active.append(["Время", "Пн"])
        for hour in range(9, 20):
            wb.active.append([f"{hour}:00", f"Занятие {hour}:00 " + "x" * 1000])

    async def model(user_content):
        # A careless model also answers for the repeated header and labels.
        events = json.loads(clock_answer(user_content))["events"]
        events.append({"cell": "B2", "title": "Занятие", "weekday": 0, "start": "09:00"})
        for event in events:
            event.setdefault("weekday", 0)
        return json.dumps({"events": events})

    with patch.object(parser, "_ask_model", side_effect=model):
        result = await run(xlsx_bytes(build))
    assert result["chunks_processed"] > 1
    assert sorted(s["start_time"] for s in result["slots"]) == [f"{hour:02d}:00" for hour in range(9, 20)]


@pytest.mark.anyio
async def test_parallel_group_schedules_ask_one_question_with_all_choices():
    def build(wb):
        wb.active.append(["День", "Время", "ПИ-241", "ПИ-242"])
        wb.active.append(["Понедельник", "09:00", "Математика", "Биология"])
    answer = json.dumps({"status": "clarification", "question": "Какую группу импортировать?",
                         "choices": ["ПИ-241", "ПИ-242"], "events": []})
    with patch.object(parser, "_ask_model", return_value=answer), \
         pytest.raises(ScheduleDocumentClarification, match="Какую группу") as error:
        await run(xlsx_bytes(build))
    assert error.value.choices == ["ПИ-241", "ПИ-242"]


@pytest.mark.anyio
async def test_user_instruction_reaches_the_model():
    def build(wb):
        wb.active.append(["День", "Время", "ПИ-241", "ПИ-242"])
        wb.active.append(["Понедельник", "09:00", "Математика", "Биология"])
    prompts = []

    async def model(user_content):
        prompts.append(user_content)
        return json.dumps({"events": [{"cell": "D2", "title": "Биология", "weekday": 0, "start": "09:00"}]})

    with patch.object(parser, "_ask_model", side_effect=model):
        result = await run(xlsx_bytes(build), "группа ПИ-242")
    assert "ИНСТРУКЦИЯ ПОЛЬЗОВАТЕЛЯ:\nгруппа ПИ-242" in prompts[0]
    assert [s["title"] for s in result["slots"]] == ["Биология"]


@pytest.mark.anyio
async def test_truncated_answer_is_retried_as_smaller_halves():
    def build(wb):
        for hour in range(8, 12):
            wb.active.append(["Пн", f"{hour}:00", f"Пара {hour}"])
    calls = []

    async def model(user_content):
        cells = fragment_cells(user_content)
        calls.append(len(cells))
        if len(cells) > 6:
            raise parser._Truncated("length")
        rows = {cell[1:] for cell, _, _ in cells}
        return json.dumps({"events": [
            {"cell": f"C{row}", "title": f"Пара {int(row) + 7}", "weekday": 0, "start": f"{int(row) + 7}:00"}
            for row in sorted(rows)
        ]})

    with patch.object(parser, "_ask_model", side_effect=model):
        result = await run(xlsx_bytes(build))
    assert calls[0] == 12
    assert len(result["slots"]) == 4


@pytest.mark.anyio
async def test_broken_json_is_retried_once_then_succeeds():
    content = xlsx_bytes(lambda wb: wb.active.append(["Понедельник", "09:00-10:00", "Тренировка"]))
    good = json.dumps({"events": [{"cell": "C1", "title": "Тренировка", "weekday": 0, "start": "09:00", "end": "10:00"}]})
    with patch.object(parser, "_ask_model", side_effect=["not json", good]) as model, \
         patch.object(parser.asyncio, "sleep", return_value=None):
        result = await run(content)
    assert model.await_count == 2
    assert result["slots"][0]["end_time"] == "10:00"


@pytest.mark.anyio
async def test_permanent_model_failure_stops_the_import():
    content = xlsx_bytes(lambda wb: wb.active.append(["Понедельник", "09:00", "Тренировка"]))
    with patch.object(parser, "_ask_model", side_effect=ValueError("empty")), \
         patch.object(parser.asyncio, "sleep", return_value=None), \
         pytest.raises(ValueError, match="Импорт остановлен"):
        await run(content)


@pytest.mark.anyio
async def test_skipped_timed_cell_is_rechecked_once():
    def build(wb):
        wb.active.append([date(2026, 9, 7), date(2026, 9, 8)])
        wb.active.append(["Йога 09:00", "Бассейн 19:00"])
    prompts = []

    async def model(user_content):
        prompts.append(user_content)
        if "ПРОВЕРКА ПРОПУСКОВ" in user_content:
            return json.dumps({"events": [
                {"cell": "B2", "title": "Бассейн", "start": "19:00"},
                {"cell": "A2", "title": "Йога", "start": "09:00"},
            ]})
        return json.dumps({"events": [{"cell": "A2", "title": "Йога", "start": "09:00"}]})

    with patch.object(parser, "_ask_model", side_effect=model):
        result = await run(xlsx_bytes(build))
    assert len(prompts) == 2
    assert "у ячеек B2 есть время" in prompts[1]
    assert [s["title"] for s in result["slots"]] == ["Йога", "Бассейн"]


@pytest.mark.anyio
async def test_formula_dates_without_cached_values_are_resolved():
    def build(wb):
        ws = wb.active
        ws.append([date(2026, 9, 7), date(2026, 9, 8)])
        ws.append(["Йога 09:00", "Бассейн 19:00"])
        ws.append(["=A1+7", "=B1+7"])
        ws.append(["Йога 09:00", "Бассейн 19:00"])
    with patch.object(parser, "_ask_model", side_effect=clock_answer):
        result = await run(xlsx_bytes(build))
    assert [s["occurrence_date"] for s in result["slots"]] == [
        "2026-09-07", "2026-09-08", "2026-09-14", "2026-09-15",
    ]


@pytest.mark.anyio
async def test_merged_label_is_carried_and_covered_cell_resolves_to_origin():
    def build(wb):
        ws = wb.active
        ws.append(["День", "Время", "Предмет"])
        for hour in range(9, 21):
            ws.append([None, f"{hour}:00", "Лекция " + "y" * 600])
        ws["A2"] = "Понедельник"
        ws.merge_cells("A2:A13")
        ws.merge_cells("C12:C13")
    seen = []

    async def model(user_content):
        seen.append(user_content)
        return json.dumps({"events": []})

    with patch.object(parser, "_ask_model", side_effect=model), pytest.raises(ValueError):
        await run(xlsx_bytes(build))
    later = [content for content in seen if "ФРАГМЕНТ: строки R2" not in content]
    assert later and all("A2 (объединено A2:A13): Понедельник" in content for content in later)

    sheet = parser._prepare_sheet(parser._read_xlsx(xlsx_bytes(build))[0])
    assert parser._resolve_cell(sheet, "C13", {13}) == (12, 3)


@pytest.mark.anyio
async def test_sheet_named_in_instruction_is_the_only_one_read_and_hidden_sheets_are_skipped():
    def build(wb):
        wb.active.title = "Группа А"
        wb.active.append(["Понедельник", "09:00", "Математика"])
        other = wb.create_sheet("Группа Б")
        other.append(["Вторник", "10:00", "Физика"])
        hidden = wb.create_sheet("Служебный")
        hidden.append(["Среда", "11:00", "Секрет"])
        hidden.sheet_state = "hidden"
    seen = []

    async def model(user_content):
        seen.append(user_content)
        return json.dumps({"events": [{"cell": "C1", "title": "Физика", "weekday": 1, "start": "10:00"}]})

    with patch.object(parser, "_ask_model", side_effect=model):
        result = await run(xlsx_bytes(build), "лист Группа Б")
    assert len(seen) == 1 and "ЛИСТ: Группа Б" in seen[0]
    assert result["verification"]["scope_labels"] == ["Группа Б"]

    seen.clear()
    with patch.object(parser, "_ask_model", side_effect=model):
        await run(xlsx_bytes(build))
    assert not any("Служебный" in content for content in seen)


@pytest.mark.anyio
async def test_dates_outside_requested_period_are_not_imported():
    def build(wb):
        wb.active.append([date(2026, 9, 7), date(2026, 12, 7)])
        wb.active.append(["Йога 09:00", "Бассейн 19:00"])
    with patch.object(parser, "_ask_model", side_effect=clock_answer):
        result = await run(xlsx_bytes(build), "", valid_range=(date(2026, 9, 1), date(2026, 9, 30)))
    assert [s["title"] for s in result["slots"]] == ["Йога"]


def test_clock_detection_ignores_counts_that_look_like_hours():
    assert parser._clocks_in("задания 1 части, в 18 ч, 9.30, 17:00") == {(18, 0), (9, 30), (17, 0)}


@pytest.mark.parametrize("start,end,deadline,expected", [
    # A deadline is the half hour before it, as for «сдать до 18:00» in chat.
    (time(23, 55), None, True, ("23:25", "23:55")),
    (time(18, 0), None, True, ("17:30", "18:00")),
    (time(0, 10), None, True, ("00:10", "00:40")),
    (time(17, 0), None, False, ("17:00", "18:30")),
    (time(23, 0), None, False, ("23:00", "23:59")),
    (time(22, 0), time(6, 0), False, ("22:00", "23:30")),
])
def test_event_interval_is_always_valid_within_one_day(start, end, deadline, expected):
    begin, finish = parser._event_interval(start, end, deadline)
    assert (begin.strftime("%H:%M"), finish.strftime("%H:%M")) == expected


def test_record_with_start_and_end_dates_is_not_a_calendar_header():
    def build(wb):
        wb.active.append(["Начало", "Конец", "Курс"])
        wb.active.append([date(2026, 9, 3), date(2026, 9, 4), "Тренинг 10:00"])
        wb.active.append([date(2026, 9, 10), date(2026, 9, 11), "Семинар 12:00"])
    sheet = parser._prepare_sheet(parser._read_xlsx(xlsx_bytes(build))[0])
    assert sheet.date_hints == {}


def multigroup_workbook(wb):
    """Notes and a timestamp above a programme → group header, as universities do."""
    ws = wb.active
    ws["C1"] = "Уважаемые обучающиеся, сверяйте расписание с сайтом"
    ws.merge_cells("C1:F1")
    ws["A2"] = "2026-10-07 23:35"
    ws["C4"], ws["E4"] = "AI360", "ЦУ"
    ws.merge_cells("C4:D4")
    ws["A6"], ws["B6"] = "Дни", "Часы"
    ws["C6"], ws["D6"], ws["E6"], ws["F6"] = "Б05-521", "Б05-522", "Б05-529", "Б05-530"
    ws["A7"], ws["B7"] = "Понедельник", "9:00 - 10:25"
    ws["C7"] = "Машинное обучение (лекция)"
    ws.merge_cells("C7:D7")
    ws["E7"], ws["F7"] = "История", "Физика"
    ws["B8"] = "10:35 - 12:00"
    ws["C8"], ws["D8"], ws["E8"] = "Физкультура", "Дискретный анализ", "Химия"


def test_header_spans_notes_timestamp_and_all_levels():
    sheet = parser._prepare_sheet(parser._read_xlsx(xlsx_bytes(multigroup_workbook))[0])
    assert parser._header_rows(sheet) == [1, 2, 4, 6]


@pytest.mark.parametrize("prompt,columns,labels", [
    ("Б05-522", {"A", "B", "C", "D"}, ["Б05-522"]),
    ("группа 521", {"A", "B", "C"}, ["Б05-521"]),
    ("AI360 б05–522", {"A", "B", "C", "D"}, ["AI360", "Б05-522"]),
    ("ЦУ", {"A", "B", "E"}, ["ЦУ"]),
    # The answer to «В AI360 несколько групп» is added to the first instruction.
    ("AI360 522", {"A", "B", "C", "D"}, ["AI360", "Б05-522"]),
])
def test_named_group_keeps_only_its_columns_and_shared_lectures(prompt, columns, labels):
    sheet = parser._prepare_sheet(parser._read_xlsx(xlsx_bytes(multigroup_workbook))[0])
    projected, found = parser._project_columns(sheet, prompt)
    data_columns = {
        parser.get_column_letter(column) for row, column in projected.texts if row >= 7
    }
    assert data_columns == columns
    assert found == labels
    # The shared lecture C7:D7 stays available to both AI360 groups.
    if "C" in columns:
        assert (7, 3) in projected.texts


def test_programme_with_two_different_groups_asks_for_the_group():
    sheet = parser._prepare_sheet(parser._read_xlsx(xlsx_bytes(multigroup_workbook))[0])
    with pytest.raises(ScheduleDocumentClarification, match="AI360") as error:
        parser._project_columns(sheet, "AI360")
    assert error.value.choices == ["Б05-521", "Б05-522"]


@pytest.mark.anyio
async def test_group_answer_reaches_the_model_without_other_groups():
    seen = []

    async def model(user_content):
        seen.append(user_content)
        return json.dumps({"events": [
            {"cell": "C7", "title": "Машинное обучение (лекция)", "weekday": 0, "start": "9:00", "end": "10:25"},
            {"cell": "D8", "title": "Дискретный анализ", "weekday": 0, "start": "10:35", "end": "12:00"},
            {"cell": "E7", "title": "История", "weekday": 0, "start": "9:00", "end": "10:25"},
        ]})

    with patch.object(parser, "_ask_model", side_effect=model):
        result = await run(xlsx_bytes(multigroup_workbook), "AI360 Б05-522")
    assert "ВЫБОР УЖЕ ВЫПОЛНЕН ПРИЛОЖЕНИЕМ: AI360, Б05-522" in seen[0]
    assert "Физкультура" not in seen[0] and "История" not in seen[0]
    assert [s["title"] for s in result["slots"]] == ["Машинное обучение (лекция)", "Дискретный анализ"]
    assert result["verification"]["scope_labels"] == ["AI360", "Б05-522"]
    assert not any("Уважаемые" in warning for warning in result["verification"]["warnings"])


@pytest.mark.anyio
async def test_question_about_an_already_named_choice_is_not_repeated():
    content = xlsx_bytes(lambda wb: wb.active.append(["Понедельник", "09:00", "Тренировка"]))
    repeated = json.dumps({"status": "clarification", "question": "Какое?", "choices": ["Тренировка"], "events": []})
    with patch.object(parser, "_ask_model", return_value=repeated), \
         pytest.raises(ValueError, match="не найдено событий"):
        await run(content, "Тренировка")


@pytest.mark.anyio
async def test_question_in_first_fragment_stops_before_reading_the_rest():
    def build(wb):
        wb.active.append(["День", "Время", "ПИ-241", "ПИ-242"])
        for hour in range(8, 20):
            wb.active.append(["Понедельник", f"{hour}:00", "Лекция " + "z" * 800, "Семинар " + "z" * 800])
    question = json.dumps({"status": "clarification", "question": "Какая группа?",
                           "choices": ["ПИ-241", "ПИ-242"], "events": []})
    with patch.object(parser, "_ask_model", return_value=question) as model, \
         pytest.raises(ScheduleDocumentClarification):
        await run(xlsx_bytes(build))
    assert model.await_count == 1


@pytest.mark.anyio
async def test_class_without_clock_in_its_cell_is_rechecked_when_skipped():
    def build(wb):
        wb.active.append(["Дни", "Часы", "Предмет"])
        wb.active.append(["Понедельник", "9:00 - 10:25", "Матанализ"])
        wb.active.append([None, "10:35 - 12:00", "Алгебра"])
    prompts = []

    async def model(user_content):
        prompts.append(user_content)
        if "ПРОВЕРКА ПРОПУСКОВ" in user_content:
            return json.dumps({"events": [{"cell": "C3", "title": "Алгебра", "weekday": 0, "start": "10:35", "end": "12:00"}]})
        return json.dumps({"events": [{"cell": "C2", "title": "Матанализ", "weekday": 0, "start": "9:00", "end": "10:25"}]})

    with patch.object(parser, "_ask_model", side_effect=model):
        result = await run(xlsx_bytes(build))
    assert "у ячеек C3 есть время" in prompts[1]
    assert [s["title"] for s in result["slots"]] == ["Матанализ", "Алгебра"]


def two_groups_sheet():
    def build(wb):
        ws = wb.active
        ws.append(["Расписание занятий, октябрь 2026"])
        ws.merge_cells("A1:D1")
        ws.append(["Дата", "Пара", "Группа ИВТ-21", "Группа ИВТ-22"])
        ws.append(["19 окт. (пн)", "1 пара 9:00–10:30", "Базы данных", "Сети ЭВМ"])
        ws.append([None, "2 пара 10:40–12:10", "Сети ЭВМ", "Базы данных"])
        ws.merge_cells("A3:A4")
    return parser._prepare_sheet(parser._read_xlsx(xlsx_bytes(build))[0])


@pytest.mark.parametrize("prompt, kept, dropped", [
    ("ИВТ-22", "D", "C"),
    ("группа ИВТ-21", "C", "D"),
    ("ивт 22", "D", "C"),  # the number belongs to one group only
    ("вторая группа", None, None),  # nothing to match: the model asks
])
def test_group_is_chosen_by_its_code_and_dates_stay(prompt, kept, dropped):
    sheet = two_groups_sheet()
    projected, labels = parser._project_columns(sheet, prompt)
    addresses = {parser._address(key) for key in projected.texts}
    if kept is None:
        assert labels == []
        return
    assert {f"{kept}3", f"{kept}4", "A3", "B3", "B4"} <= addresses
    assert not {f"{dropped}3", f"{dropped}4"} & addresses


def test_date_with_weekday_is_an_axis_label():
    assert parser._is_axis_label("19 окт. (пн)")
    assert parser._is_axis_label("Пн, 12.10")
    assert not parser._is_axis_label("19 окт. экзамен")


def test_weekday_of_a_day_block_is_computed_not_guessed():
    def build(wb):
        ws = wb.active
        ws.append(["Дни", "Часы", "Группа 1"])
        ws.append(["Понедельник", "9:00 - 10:25", "Алгебра"])
        ws.append([None, "10:35 - 12:00", "Физика"])
        ws.append(["Вторник", "9:00 - 10:25", "Химия"])
        ws.append([None, "10:35 - 12:00", None])
        ws.merge_cells("A2:A3")
        ws.merge_cells("A4:A5")
    sheet = parser._prepare_sheet(parser._read_xlsx(xlsx_bytes(build))[0])
    assert {parser._address(key): day for key, day in sheet.weekday_hints.items() if key[1] == 3} == {
        "C2": 0, "C3": 0, "C4": 1,
    }
    chunk = parser.Chunk(sheet, sheet.rows)
    # The model says Thursday; the label beside the class says Tuesday.
    slot, _ = parser._build_slot(chunk, {"cell": "C4", "title": "Химия", "weekday": 3, "start": "9:00"}, None)
    assert slot["weekday"] == 1 and slot["occurrence_date"] is None
    assert "C4 {вт}: Химия" in parser._render_chunk(chunk, len(sheet.rows))


def test_a_date_the_table_never_mentions_is_ignored():
    def build(wb):
        wb.active.append(["Понедельник", "9:00", "Алгебра"])
        wb.active.append(["Вторник", "9:00", "Химия"])
    sheet = parser._prepare_sheet(parser._read_xlsx(xlsx_bytes(build))[0])
    chunk = parser.Chunk(sheet, sheet.rows)
    slot, _ = parser._build_slot(chunk, {"cell": "C1", "title": "Алгебра", "date": "2026-10-08", "weekday": 0, "start": "9:00"}, None)
    assert slot["occurrence_date"] is None and slot["weekday"] == 0


def mirea_sheet():
    """Two groups side by side with «9-00» times and a week I / II column."""
    def build(wb):
        ws = wb.active
        ws.append(["День недели", "№ пары", "Нач. занятий", "Оконч. занятий", "Неделя", "КМБО-01-19", "19", "КМБО-03-19"])
        ws.append(["ПОНЕДЕЛЬНИК", 1, "9-00", "10-30", "I", "кр. 1 н. Физика", "СР", "Химия"])
        ws.append([None, None, None, None, "II", "Алгебра", "СР", "Химия"])
        ws.append([None, 2, "10-40", "12-10", "I", "Военная подготовка", None, None])
        ws.append([None, None, None, None, "II", None, None, None])
        ws.append(["ВТОРНИК", 1, "9-00", "10-30", "I", "История", None, None])
        ws.merge_cells("A2:A5"); ws.merge_cells("B2:B3"); ws.merge_cells("C2:C3"); ws.merge_cells("D2:D3")
        ws.merge_cells("B4:B5"); ws.merge_cells("C4:C5"); ws.merge_cells("D4:D5"); ws.merge_cells("F4:F5")
    return parser._prepare_sheet(parser._read_xlsx(xlsx_bytes(build))[0])


def test_dash_times_week_column_and_weekdays_are_computed():
    sheet = mirea_sheet()
    assert sheet.texts[(2, 3)] == "9:00" and (10, 40) in sheet.clocks
    # «СР» (самостоятельная работа) repeated in a column is not Wednesday.
    assert sheet.weekday_hints[(3, 6)] == 0 and sheet.weekday_hints[(6, 6)] == 1
    assert sheet.week_hints[(2, 6)] == "odd" and sheet.week_hints[(3, 6)] == "even"
    assert sheet.week_hints[(4, 6)] == "every"  # merged over weeks I and II


def test_class_numbers_are_not_dash_times():
    def build(wb):
        ws = wb.active
        ws.append(["Дни", "Часы", "Группа"])
        for hours in ("1-2", "3-4", "9-10"):
            ws.append(["понедельник", hours, "Математика"])
    sheet = parser._prepare_sheet(parser._read_xlsx(xlsx_bytes(build))[0])
    assert sheet.clocks == set()


def test_part_of_a_group_code_selects_nothing_else():
    projected, labels = parser._project_columns(mirea_sheet(), "КМБО-01-19")
    assert labels == ["КМБО-01-19"]
    assert not any(key[1] in (7, 8) for key in projected.texts)
    assert parser._group_codes([mirea_sheet()]) == ["КМБО-01-19", "КМБО-03-19"]


def test_week_mark_of_the_cell_stays_with_the_title():
    assert parser._with_week_mark("Кр. Физика", "кр. 1 н. Физика") == "кр. 1 н. Физика"


def test_repeated_day_column_keeps_the_time_column():
    """«Дни | Часы | группы… | Дни | группы»: the times stay with any group."""
    def build(wb):
        ws = wb.active
        ws.append(["Дни", "Часы", "Б05-511", "Дни", "Б05-522"])
        ws.append(["Понедельник", "9:00 - 10:25", "Алгебра", "Понедельник", "Физика"])
        ws.append(["Вторник", "10:35 - 12:00", "Химия", "Вторник", "История"])
    sheet = parser._prepare_sheet(parser._read_xlsx(xlsx_bytes(build))[0])
    projected, _ = parser._project_columns(sheet, "522")
    columns = {key[1] for key in projected.texts}
    assert {2, 4, 5} <= columns and 3 not in columns
