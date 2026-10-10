"""Test spreadsheets with a known ground truth (truth.json)."""
import json
from datetime import date, datetime, time
from openpyxl import Workbook
from openpyxl.styles import Font

truth = {}

# 1. Plain list: date | time | subject | room
wb = Workbook(); ws = wb.active; ws.title = "Расписание"
ws.append(["Дата", "Время", "Предмет", "Аудитория"])
rows = [
    (date(2026, 10, 12), "09:00-10:30", "Математический анализ", "ауд. 301"),
    (date(2026, 10, 12), "10:45-12:15", "Линейная алгебра", "ауд. 214"),
    (date(2026, 10, 13), "13:00-14:30", "История России", "ауд. 105"),
    (date(2026, 10, 14), "09:00-10:30", "Физика", "лаб. 2"),
    (date(2026, 10, 15), "15:00-16:30", "Английский язык", "ауд. 410"),
    (date(2026, 10, 16), "11:00-12:30", "Программирование", "ауд. 501"),
]
for r in rows:
    ws.append(list(r))
wb.save("xlsx/1_list.xlsx")
truth["1_list.xlsx"] = [(r[0].isoformat(), r[1][:5], r[1][6:], r[2]) for r in rows]

# 2. Weekly grid without dates (needs a period)
wb = Workbook(); ws = wb.active; ws.title = "Неделя"
ws.append(["Время", "Понедельник", "Вторник", "Среда", "Четверг", "Пятница"])
ws.append(["8:30–10:00", "Алгебра", "", "Химия", "", "Физкультура"])
ws.append(["10:10–11:40", "", "Литература", "", "Биология", ""])
ws.append(["11:50–13:20", "Информатика", "", "", "", "Обществознание"])
wb.save("xlsx/2_weekly.xlsx")
truth["2_weekly.xlsx"] = {"weekly": [
    (0, "08:30", "10:00", "Алгебра"), (2, "08:30", "10:00", "Химия"), (4, "08:30", "10:00", "Физкультура"),
    (1, "10:10", "11:40", "Литература"), (3, "10:10", "11:40", "Биология"),
    (0, "11:50", "13:20", "Информатика"), (4, "11:50", "13:20", "Обществознание")]}

# 3. Two groups side by side, merged date column, Russian date text
wb = Workbook(); ws = wb.active; ws.title = "Октябрь"
ws.append(["Расписание занятий, октябрь 2026"])
ws.merge_cells("A1:D1")
ws.append(["Дата", "Пара", "Группа ИВТ-21", "Группа ИВТ-22"])
data3 = [
    ("19 окт. (пн)", "1 пара 9:00–10:30", "Базы данных (Петров)", "Сети ЭВМ"),
    (None, "2 пара 10:40–12:10", "Сети ЭВМ", "Базы данных"),
    ("20 окт. (вт)", "1 пара 9:00–10:30", "", "Операционные системы"),
    (None, "2 пара 10:40–12:10", "Операционные системы", ""),
    ("21 окт. (ср)", "3 пара 12:40–14:10", "Физкультура", "Физкультура"),
]
for r in data3:
    ws.append(list(r))
ws.merge_cells("A3:A4"); ws.merge_cells("A5:A6")
wb.save("xlsx/3_two_groups.xlsx")
truth["3_two_groups.xlsx"] = {"ИВТ-21": [
    ("2026-10-19", "09:00", "10:30", "Базы данных"), ("2026-10-19", "10:40", "12:10", "Сети ЭВМ"),
    ("2026-10-20", "10:40", "12:10", "Операционные системы"), ("2026-10-21", "12:40", "14:10", "Физкультура")],
    "ИВТ-22": [("2026-10-19", "09:00", "10:30", "Сети ЭВМ"), ("2026-10-19", "10:40", "12:10", "Базы данных"),
    ("2026-10-20", "09:00", "10:30", "Операционные системы"), ("2026-10-21", "12:40", "14:10", "Физкультура")]}

# 4. Personal plan with deadlines, mixed formats, a note row and an empty sheet
wb = Workbook(); ws = wb.active; ws.title = "План"
ws.append(["Что", "Когда", "Комментарий"])
ws.append(["Сдать лабораторную №3", "до 23.10.2026 18:00", "в Moodle"])
ws.append(["Встреча с научруком", datetime(2026, 10, 22, 16, 0), "каб. 12"])
ws.append(["Олимпиада по программированию", "25.10.2026 10:00–14:00", ""])
ws.append(["Примечание: расписание может меняться", None, None])
wb.create_sheet("Пусто")
wb.save("xlsx/4_deadlines.xlsx")
truth["4_deadlines.xlsx"] = [("2026-10-23", None, "18:00", "лабораторн"), ("2026-10-22", "16:00", None, "научрук"),
                             ("2026-10-25", "10:00", "14:00", "Олимпиад")]

# 5. Calendar grid: week columns with dates in header, times in rows
wb = Workbook(); ws = wb.active; ws.title = "Курс"
ws.append(["", datetime(2026, 11, 2), datetime(2026, 11, 3), datetime(2026, 11, 4), datetime(2026, 11, 5)])
ws.append(["", "пн", "вт", "ср", "чт"])
ws.append([time(18, 0), "Вебинар: Право", "", "Вебинар: Экономика", ""])
ws.append([time(19, 30), "", "Разбор ДЗ", "", "Пробный ЕГЭ"])
for c in ws[1][1:]:
    c.number_format = "DD.MM.YYYY"
wb.save("xlsx/5_calendar_grid.xlsx")
truth["5_calendar_grid.xlsx"] = [("2026-11-02", "18:00", None, "Право"), ("2026-11-04", "18:00", None, "Экономика"),
                                 ("2026-11-03", "19:30", None, "Разбор"), ("2026-11-05", "19:30", None, "Пробный")]
json.dump(truth, open("xlsx/truth.json", "w"), ensure_ascii=False, indent=1)
print("ok")
